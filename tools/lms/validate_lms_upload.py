#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
StudyPeak 수학 LMS 업로드 파일 검증기 (validate_lms_upload.py)
================================================================
빌드가 끝난 ExamTemplate/TextbookTemplate xlsx가 실제로 LMS에 오류 없이
올라갈 수 있는 상태인지 "사람이 눈으로 훑어보기 전에" 기계적으로 확인한다.

사용법:
    python3 validate_lms_upload.py <검증할_파일.xlsx> [--master 마스터유형표.xlsx] [--sheet ExamTemplate|TextbookTemplate]

종료 코드:
    0 = 모든 항목 PASS (전달 가능)
    1 = 하나 이상 ERROR (전달 금지 — 고쳐서 재검증)

이 스크립트는 "완전 자동 수정"을 하지 않는다. 실패 항목을 정확한 행/열과 함께
보고하는 것까지가 이 스크립트의 역할이며, 수정은 빌드 스크립트(Claude)가 하고
다시 이 스크립트를 돌려서 통과할 때까지 반복한다.

변경 이력
- 2026-09-30: LMS 업데이트(exam_template 2.xlsx, v3) 반영. 그룹 헤더 8·9열 라벨이
  "Page From/Page To" → "시작 페이지/종료 페이지"로 바뀜. v3를 표준으로 삼고,
  구 v2 영문/한글 라벨은 WARNING으로 낮춤. 단답형 답안개수 필수 규칙은 유지.
- 2026-10-29: load_master 버그 수정. openpyxl의 max_row는 서식만 남은 빈 꼬리 행까지
  포함하는데, 5열 전체를 ffill 하면 마스터 끝의 카테고리 선언 행(추론/실생활 적용/
  자료 해석)과 빈 행들이 직전 (소분류,유형)을 그대로 물려받아, 고등 마스터 기준
  유형 734개가 1,682개로 부풀고 대분류가 뒤섞였다. 이제 유형(E열) 원본이 빈 행은 건너뛴다.
"""
import sys
import re
import argparse
import zipfile
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl이 필요합니다: pip install openpyxl --break-system-packages")
    sys.exit(2)

import warnings
warnings.filterwarnings("ignore")

# ---------------------------------------------------------------------------
# 라벨 상수
# ---------------------------------------------------------------------------
GROUP_HEADER_LABELS = {"Group Order", "그룹 순서"}
ITEM_HEADER_LABELS = {"Number of questions", "문항 번호", "문항번호"}

# ★ v3 (2026-09-30 LMS 업데이트, exam_template 2.xlsx) — 현재 표준
GROUP_HEADER_EXPECTED_V3 = ["그룹 순서", "대분류", "소분류", "유형",
                            "지문 제목", "지문 작성자", "지문 태그", "시작 페이지", "종료 페이지"]
ITEM_HEADER_EXPECTED_V3 = [
    "문항 번호", "대분류", "소분류", "유형", "정답유형", "난이도", "권장풀이시간(mm:ss)",
    "답안개수", "점수", "정답선택지", "주관식 정답", "비교 유형", "단위", "답안없음",
]

# 구 v2 (영문 / 한글 Page From·To) — 폐기됨, WARNING
GROUP_HEADER_EXPECTED_EN = ["Group Order", "Category", "Subcategory", "Question Type",
                            "Article Title", "Article Author", "Article Tag", "Page From", "Page To"]
GROUP_HEADER_EXPECTED_KO_V2 = ["그룹 순서", "대분류", "소분류", "유형",
                               "지문 제목", "지문 작성자", "지문 태그", "Page From", "Page To"]
GROUP_HEADER_EXPECTED_KO_V1 = ["그룹 순서", "대분류", "소분류", "유형", "지문 제목", "작성자"]

ITEM_HEADER_EXPECTED_EN_V2 = [
    "Number of questions", "Category", "Subcategory", "Question Type", "Question Answer Type",
    "Level", "Recommend AnswerTime (mm:ss)", "Number of answers", "Score",
    "Correct Choices", "Correct Textual Answers", "Compare Type", "Unit", "Without Answers",
]
ITEM_HEADER_EXPECTED_KO_V1 = [
    "문항번호", "대분류", "소분류", "유형", "정답유형", "난이도", "권장풀이시간(mm:ss)",
    "답안개수", "점수", "정답선택지", "정답단답형", "단위", "답안없음",
]

CHOICE_TYPES = {"Single Choice", "Multiple Choice", "단일 선택", "다중 선택", "5지선다형"}
SHORT_TYPES = {"Short Answer", "단답형"}

CATEGORY_10 = {"함수", "수와 연산", "방정식과 부등식", "논리", "기하", "통계",
               "경우의 수와 확률", "추론", "자료 해석", "실생활 적용"}
GRADE_7 = {"공통수학1", "공통수학2", "대수", "미적분1", "미적분2", "확률과 통계", "기하"}

ANSWER_TYPE_OK = {"단일 선택", "다중 선택", "단답형"}
LEVEL_OK = {"일반적인 문제", "쉬운 문제", "변별력 문제", "함정 문제", "어려운 문제", "엄청 어려운 문제"}
NOANS_OK = {"아니오", "예"}
COMPARE_OK = {"순서 있음", "순서 없음", "유의어 처리"}


class Report:
    def __init__(self):
        self.errors = []
        self.warnings = []
        self.passed = []

    def err(self, code, msg):
        self.errors.append((code, msg))

    def warn(self, code, msg):
        self.warnings.append((code, msg))

    def ok(self, code, msg):
        self.passed.append((code, msg))

    @staticmethod
    def _print_grouped(items, symbol, cap=5):
        by_code = {}
        for code, msg in items:
            by_code.setdefault(code, []).append(msg)
        for code, msgs in by_code.items():
            shown = msgs if len(msgs) <= cap else msgs[:cap]
            for m in shown:
                print(f"  {symbol} {code}: {m}")
            if len(msgs) > cap:
                print(f"  {symbol} {code}: ... 외 {len(msgs) - cap}건 (총 {len(msgs)}건, 같은 원인일 가능성 높음)")

    def print_report(self, filename):
        print("=" * 70)
        print(f"검증 대상: {filename}")
        print("=" * 70)
        if self.passed:
            print(f"\n[PASS] {len(self.passed)}건")
            self._print_grouped(self.passed, "OK")
        if self.warnings:
            print(f"\n[WARNING] {len(self.warnings)}건 (전달은 가능하나 확인 권장)")
            self._print_grouped(self.warnings, "!!")
        if self.errors:
            print(f"\n[ERROR] {len(self.errors)}건 (전달 금지 — 수정 후 재검증 필요)")
            self._print_grouped(self.errors, "XX")
        print("\n" + "=" * 70)
        if self.errors:
            print(f"최종 판정: FAIL — ERROR {len(self.errors)}건 해결 전까지 사용자에게 전달하지 말 것")
        else:
            print("최종 판정: PASS — 전달 가능")
        print("=" * 70)


def load_master(master_path):
    """마스터 유형표 로딩 → (sub,qt) 유효셋과 (sub,qt)->category 역조회 dict 반환."""
    df_rows = []
    wb = openpyxl.load_workbook(master_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    last = [None, None, None, None, None]
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, 6)]
        raw_qt = vals[4]                     # ★ 2026-10-29: ffill 전 원본
        for i in range(5):
            if vals[i] not in (None, ""):
                last[i] = vals[i]
            else:
                vals[i] = last[i]
        if raw_qt in (None, ""):             # ★ 유형이 원래 비어 있으면 데이터 행이 아님
            continue
        cat, g1, sub, grade, qt = vals
        if sub and qt:
            df_rows.append((cat, grade, sub, qt))
    valid_pairs = set((sub, qt) for _, _, sub, qt in df_rows)
    cat_lookup = {}
    conflicts = []
    for cat, grade, sub, qt in df_rows:
        key = (sub, qt)
        if key in cat_lookup and cat_lookup[key] != cat:
            conflicts.append((key, cat_lookup[key], cat))
        cat_lookup[key] = cat
    return valid_pairs, cat_lookup, conflicts


def classify_rows(ws, max_col=14):
    rows = []
    state = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        vals = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        a_str = str(a).strip() if a is not None else ""
        if a_str in GROUP_HEADER_LABELS:
            rows.append((r, "group_header", vals))
            state = "after_group_header"
            continue
        if state == "after_group_header":
            rows.append((r, "group_data", vals))
            state = None
            continue
        if a_str in ITEM_HEADER_LABELS:
            rows.append((r, "item_header", vals))
            state = "in_item_block"
            continue
        if state == "in_item_block" and any(v not in (None, "") for v in vals):
            rows.append((r, "item_row", vals))
            continue
    return rows


def check_header_orders(rows, report):
    n_g3 = n_i3 = 0
    for r, kind, vals in rows:
        if kind == "group_header":
            v9 = [str(x) if x is not None else "" for x in vals[:9]]
            if v9 == GROUP_HEADER_EXPECTED_V3:
                n_g3 += 1
            elif v9 in (GROUP_HEADER_EXPECTED_EN, GROUP_HEADER_EXPECTED_KO_V2):
                report.warn("HEADER-GROUP-V2",
                            f"{r}행: 구 v2 그룹 헤더(Page From/To 또는 영문 라벨) — v3 표준은 {GROUP_HEADER_EXPECTED_V3}")
            elif v9[:6] == GROUP_HEADER_EXPECTED_KO_V1:
                report.warn("HEADER-GROUP-V1",
                            f"{r}행: 구버전(v1, 6열) 그룹 헤더 — v3(9열)로 마이그레이션 필요")
            else:
                report.err("HEADER-GROUP",
                           f"{r}행 그룹 헤더 열 순서 불일치: {v9} (기대: {GROUP_HEADER_EXPECTED_V3})")
        if kind == "item_header":
            v14 = [str(x) if x is not None else "" for x in vals[:14]]
            if v14 == ITEM_HEADER_EXPECTED_V3:
                n_i3 += 1
            elif v14 == ITEM_HEADER_EXPECTED_EN_V2:
                report.warn("HEADER-ITEM-V2",
                            f"{r}행: 문항 헤더가 영문 라벨(v2) — v3 표준은 한글 라벨")
            elif v14[:13] == ITEM_HEADER_EXPECTED_KO_V1:
                report.warn("HEADER-ITEM-V1",
                            f"{r}행: 구버전(v1, 13열) 문항 헤더 — v3(14열)로 재작업 필요")
            else:
                report.err("HEADER-ITEM", f"{r}행 문항 헤더 열 순서 불일치: {v14}")
    if n_g3:
        report.ok("HEADER-GROUP", f"그룹 헤더 {n_g3}개 전부 v3 규격(한글 9열) 일치")
    if n_i3:
        report.ok("HEADER-ITEM", f"문항 헤더 {n_i3}개 전부 v3 규격(한글 14열) 일치")


def check_master_consistency(rows, valid_pairs, cat_lookup, report):
    checked = 0
    for r, kind, vals in rows:
        if kind not in ("group_data", "item_row"):
            continue
        b, c, d = vals[1], vals[2], vals[3]
        if not (b and c and d):
            continue
        checked += 1
        key = (str(c).strip(), str(d).strip())
        if key not in valid_pairs:
            report.err("MASTER-MISMATCH",
                       f"{r}행: (소분류='{c}', 유형='{d}') 조합이 마스터 유형표에 없음")
            continue
        expected_cat = cat_lookup[key]
        if str(b).strip() != expected_cat:
            if str(b).strip() in GRADE_7:
                report.err("CATEGORY-IS-GRADE",
                           f"{r}행: 대분류 칸에 과목명('{b}')이 들어감 (정답: '{expected_cat}')")
            else:
                report.err("CATEGORY-MISMATCH",
                           f"{r}행: 대분류 불일치 — 파일='{b}', 정답='{expected_cat}'")
        elif str(b).strip() not in CATEGORY_10:
            report.warn("CATEGORY-UNKNOWN", f"{r}행: 대분류 '{b}'가 확정 10종 목록에 없음")
    if checked == 0:
        report.warn("MASTER-CHECK-EMPTY", "분류가 채워진 행을 찾지 못함")
    else:
        n_err = sum(1 for c, _ in report.errors
                    if c in ("MASTER-MISMATCH", "CATEGORY-MISMATCH", "CATEGORY-IS-GRADE"))
        if n_err == 0:
            report.ok("MASTER-CHECK", f"분류 채워진 {checked}개 행 전부 마스터와 일치")


def check_answers(rows, report):
    item_rows = [(r, vals) for r, kind, vals in rows if kind == "item_row"]
    clusters = []
    cur_key, cur = None, []
    for r, vals in item_rows:
        a = str(vals[0]).strip() if vals[0] is not None else ""
        if a != cur_key:
            if cur:
                clusters.append((cur_key, cur))
            cur_key, cur = a, [(r, vals)]
        else:
            cur.append((r, vals))
    if cur:
        clusters.append((cur_key, cur))

    for key, cluster_rows in clusters:
        primary = None
        for r, vals in cluster_rows:
            if vals[4] not in (None, ""):
                primary = (r, vals)
                break
        if primary is None:
            continue
        pr, pvals = primary
        answer_type = str(pvals[4]).strip()
        if pvals[7] in (None, ""):
            report.err("MISSING-NUM-ANSWERS",
                       f"{pr}행 (문항 {key}): 답안개수가 비어 있음 — 단답형도 예외 아님")
        if pvals[5] not in (None, "") and str(pvals[5]).strip() not in LEVEL_OK:
            report.err("BAD-LEVEL", f"{pr}행 (문항 {key}): 난이도 '{pvals[5]}'는 _Lookup 허용값이 아님")
        if pvals[13] not in (None, "") and str(pvals[13]).strip() not in NOANS_OK:
            report.err("BAD-NOANSWER", f"{pr}행 (문항 {key}): 답안없음 '{pvals[13]}'는 아니오/예가 아님")
        if pvals[11] not in (None, "") and str(pvals[11]).strip() not in COMPARE_OK:
            report.err("BAD-COMPARE", f"{pr}행 (문항 {key}): 비교 유형 '{pvals[11]}' 허용값 아님")
        all_j = [str(v[9]).strip() for _, v in cluster_rows if v[9] not in (None, "")]
        all_k = [str(v[10]).strip() for _, v in cluster_rows if v[10] not in (None, "")]
        if answer_type in CHOICE_TYPES:
            if not all_j:
                report.err("MISSING-CHOICE-ANSWER", f"{pr}행 (문항 {key}): 정답선택지 비어 있음")
            if any("<" in x for x in all_j):
                report.err("HTML-WRAPPED", f"{pr}행 (문항 {key}): 정답에 HTML 태그 포함")
        elif answer_type in SHORT_TYPES:
            if not all_k:
                report.err("MISSING-TEXT-ANSWER", f"{pr}행 (문항 {key}): 주관식 정답 비어 있음")
            if any("<" in x for x in all_k):
                report.err("HTML-WRAPPED", f"{pr}행 (문항 {key}): 주관식 정답에 HTML 태그 포함")
        else:
            report.warn("UNKNOWN-ANSWER-TYPE", f"{pr}행 (문항 {key}): 정답유형 '{answer_type}' 미인식")

    n_err = sum(1 for c, _ in report.errors if c.startswith("MISSING-") or c.startswith("BAD-"))
    if item_rows and n_err == 0:
        report.ok("ANSWER-CHECK", f"문항 클러스터 {len(clusters)}개 전부 정답/답안개수/열거값 정상")


def check_time_and_score(rows, ws, report, expect_total_sec=None, expect_total_score=None):
    """권장풀이시간 합계·배점 합계를 집계해 보고한다."""
    import datetime
    tot_sec, tot_score, n = 0, 0, 0
    for r, kind, vals in rows:
        if kind != "item_row" or vals[4] in (None, ""):
            continue
        n += 1
        g = vals[6]
        if isinstance(g, datetime.time):
            tot_sec += g.hour * 3600 + g.minute * 60 + g.second
        else:
            report.err("BAD-TIME", f"{r}행: 권장풀이시간이 time 객체가 아님 ({g!r})")
        if ws.cell(row=r, column=7).number_format != "mm:ss":
            report.warn("TIME-FORMAT", f"{r}행: 권장풀이시간 서식이 mm:ss가 아님")
        if isinstance(vals[8], (int, float)):
            tot_score += vals[8]
        else:
            report.err("BAD-SCORE", f"{r}행: 점수가 숫자가 아님 ({vals[8]!r})")
    tot_score = round(tot_score, 6)   # 0.01 단위 배점의 부동소수점 누적오차 제거
    if tot_score == int(tot_score):
        tot_score = int(tot_score)
    msg = f"문항 {n}개 / 권장풀이시간 합계 {tot_sec}초({tot_sec/60:.1f}분) / 배점 합계 {tot_score}점"
    if expect_total_sec is not None and tot_sec != expect_total_sec:
        report.err("TIME-SUM", f"{msg} — 기대 {expect_total_sec}초와 불일치")
    elif expect_total_score is not None and abs(tot_score - expect_total_score) > 1e-6:
        report.err("SCORE-SUM", f"{msg} — 기대 {expect_total_score}점과 불일치")
    else:
        report.ok("TIME-SCORE-SUM", msg)


def check_number_format(ws, rows, report):
    bad = []
    for r, kind, vals in rows:
        if kind != "item_row":
            continue
        cell = ws.cell(row=r, column=1)
        if cell.value not in (None, "") and cell.number_format != "@":
            bad.append(r)
    if bad:
        report.warn("NUMBER-FORMAT", f"문항 번호 열(A) 서식이 '@'가 아닌 행: {bad[:10]}")
    else:
        report.ok("NUMBER-FORMAT", "문항 번호 열 전부 텍스트 서식('@') 확인")


def check_self_closing_cells(path, report):
    total, examples = 0, []
    with zipfile.ZipFile(path) as z:
        for name in [n for n in z.namelist()
                     if n.startswith("xl/worksheets/") and n.endswith(".xml")]:
            xml = z.read(name).decode("utf-8", errors="ignore")
            found = re.findall(r'<c\s+r="[A-Z]+\d+"[^>/]*/>', xml)
            total += len(found)
            if found and len(examples) < 3:
                examples.extend(found[:3 - len(examples)])
    if total > 0:
        report.err("SELF-CLOSING-CELLS",
                   f"자기닫힘 <c r=.../> 빈 셀 태그 {total}건 (예: {examples[:3]})")
    else:
        report.ok("SELF-CLOSING-CELLS", "자기닫힘 빈 셀 태그 0건")


def check_lookup(wb, report):
    if "_Lookup" not in wb.sheetnames:
        report.err("LOOKUP-MISSING", "_Lookup 시트가 없음")
        return
    ws = wb["_Lookup"]
    if ws.sheet_state != "hidden":
        report.warn("LOOKUP-VISIBLE", "_Lookup 시트가 숨김 상태가 아님")
    else:
        report.ok("LOOKUP-SHEET", "_Lookup 시트 존재 및 숨김 상태 확인")


def check_filename(path, report):
    name = Path(path).name
    try:
        name.encode("ascii")
        report.ok("FILENAME-ASCII", f"파일명 ASCII 확인: {name}")
    except UnicodeEncodeError:
        report.err("FILENAME-NON-ASCII", f"파일명에 비-ASCII 문자 포함: {name}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("target")
    ap.add_argument("--master")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--expect-sec", type=int, default=None)
    ap.add_argument("--expect-score", type=float, default=None)
    args = ap.parse_args()

    report = Report()
    check_filename(args.target, report)

    wb = openpyxl.load_workbook(args.target, data_only=False)
    sheet_name = args.sheet
    if sheet_name is None:
        for cand in ("ExamTemplate", "TextbookTemplate"):
            if cand in wb.sheetnames:
                sheet_name = cand
                break
    if sheet_name is None or sheet_name not in wb.sheetnames:
        report.err("SHEET-NOT-FOUND", f"시트를 찾을 수 없음 (실제: {wb.sheetnames})")
        report.print_report(args.target)
        sys.exit(1)

    ws = wb[sheet_name]
    rows = classify_rows(ws)

    check_header_orders(rows, report)
    if args.master:
        valid_pairs, cat_lookup, conflicts = load_master(args.master)
        for key, c1, c2 in conflicts[:5]:
            report.warn("MASTER-INTERNAL-CONFLICT",
                        f"마스터 내 (소분류,유형)={key}에 대분류 2개: '{c1}' vs '{c2}'")
        check_master_consistency(rows, valid_pairs, cat_lookup, report)
    else:
        report.warn("MASTER-SKIPPED", "--master 없이 실행 — 마스터 일치성 검사 생략")

    check_answers(rows, report)
    check_time_and_score(rows, ws, report, args.expect_sec, args.expect_score)
    check_number_format(ws, rows, report)
    check_self_closing_cells(args.target, report)
    check_lookup(wb, report)

    report.print_report(args.target)
    sys.exit(1 if report.errors else 0)


if __name__ == "__main__":
    main()
