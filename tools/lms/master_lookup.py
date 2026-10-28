#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
마스터 유형표 로딩·역조회 (master_lookup.py)
=============================================
스터디피크 수학 마스터 유형표(5열, 셀 병합 → ffill 필요)를 읽어
  · valid3   : (과목, 소분류, 유형) 유효 조합 집합
  · cat3     : (과목, 소분류, 유형) → 대분류
  · cat2     : (소분류, 유형)       → 대분류   (과목 미상일 때 폴백)
를 돌려준다. 결과는 pickle로 캐시한다.

**마스터 문자열은 절대 수정하지 않는다.** 오타처럼 보여도 그대로 쓴다
(∑, 콜론, 전각문자 ｏ 등 특수문자 보존).

사용 (모듈):
    from master_lookup import load
    m = load('data/스터디피크_수학_고등_유형.xlsx')
    m.category('미적분2', '여러 가지 미분법', '합성함수의 미분법')

사용 (CLI):
    python3 master_lookup.py data/스터디피크_수학_고등_유형.xlsx            # 통계
    python3 master_lookup.py data/스터디피크_수학_고등_유형.xlsx --grade 대수
    python3 master_lookup.py data/스터디피크_수학_고등_유형.xlsx --find 미분
"""
import argparse
import pickle
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl이 필요합니다: pip install openpyxl --break-system-packages")
    sys.exit(2)

CACHE_VERSION = 2   # 로딩 로직이 바뀌면 올릴 것 (캐시 무효화)

CATEGORY_10 = ["함수", "수와 연산", "방정식과 부등식", "논리", "기하", "통계",
               "경우의 수와 확률", "추론", "자료 해석", "실생활 적용"]
GRADE_7 = ["공통수학1", "공통수학2", "대수", "미적분1", "미적분2", "확률과 통계", "기하"]


class Master:
    def __init__(self, rows):
        # rows: list of (cat, grade, sub, qt)
        self.rows = rows
        self.valid3 = {(g, s, q) for _, g, s, q in rows}
        self.valid2 = {(s, q) for _, _, s, q in rows}
        self.cat3 = {(g, s, q): c for c, g, s, q in rows}
        self.cat2 = {}
        self.conflicts = []
        for c, g, s, q in rows:
            key = (s, q)
            if key in self.cat2 and self.cat2[key] != c:
                self.conflicts.append((key, self.cat2[key], c))
            self.cat2[key] = c

    def category(self, grade, sub, qt):
        """(과목, 소분류, 유형) → 대분류. 없으면 None."""
        if (grade, sub, qt) in self.cat3:
            return self.cat3[(grade, sub, qt)]
        return self.cat2.get((sub, qt))

    def is_valid(self, grade, sub, qt):
        return (grade, sub, qt) in self.valid3

    def check_all(self, triples):
        """[(grade, sub, qt), ...] 중 마스터에 없는 것만 돌려준다."""
        return [t for t in triples if t not in self.valid3]

    def grades(self):
        return sorted({g for _, g, _, _ in self.rows}, key=str)

    def types_of(self, grade=None, sub=None):
        out = []
        for c, g, s, q in self.rows:
            if grade is not None and g != grade:
                continue
            if sub is not None and s != sub:
                continue
            out.append((c, g, s, q))
        return out

    def find(self, needle):
        n = str(needle)
        return [r for r in self.rows if n in str(r[2]) or n in str(r[3])]


def _read(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows, last = [], [None] * 5
    for r in range(2, ws.max_row + 1):          # 1행은 헤더
        vals = [ws.cell(row=r, column=c).value for c in range(1, 6)]
        raw_qt = vals[4]                        # ★ ffill 전 원본
        for i in range(5):                      # 병합 셀 ffill
            if vals[i] not in (None, ""):
                last[i] = vals[i]
            else:
                vals[i] = last[i]
        # ★ 유형(E열) 원본이 비어 있으면 데이터 행이 아니다.
        #   마스터 끝의 카테고리 선언 행(추론/실생활 적용/자료 해석)과 빈 꼬리 행까지
        #   ffill 해 버리면 직전 (소분류,유형)이 엉뚱한 대분류로 수백 건 복제된다.
        if raw_qt in (None, ""):
            continue
        cat, _g1, sub, grade, qt = vals
        if sub and qt:
            rows.append((cat, grade, sub, qt))
    return rows


def load(path, use_cache=True):
    """마스터 유형표를 읽어 Master 객체를 돌려준다. mtime 기준으로 캐시."""
    path = Path(path)
    cache = path.with_suffix(".cache.pkl")
    if use_cache and cache.exists() and cache.stat().st_mtime >= path.stat().st_mtime:
        try:
            with open(cache, "rb") as f:
                ver, rows = pickle.load(f)
            if ver == CACHE_VERSION:
                return Master(rows)
        except Exception:
            pass
    rows = _read(path)
    if use_cache:
        try:
            with open(cache, "wb") as f:
                pickle.dump((CACHE_VERSION, rows), f)
        except Exception:
            pass
    return Master(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("master")
    ap.add_argument("--grade", help="해당 과목의 (소분류, 유형)만 출력")
    ap.add_argument("--sub", help="해당 소분류만 출력")
    ap.add_argument("--find", help="소분류·유형 문자열 부분 검색")
    ap.add_argument("--no-cache", action="store_true")
    args = ap.parse_args()

    m = load(args.master, use_cache=not args.no_cache)

    if args.find:
        hits = m.find(args.find)
        for c, g, s, q in hits:
            print(f"{c}\t{g}\t{s}\t{q}")
        print(f"\n— {len(hits)}건")
        return

    if args.grade or args.sub:
        hits = m.types_of(args.grade, args.sub)
        for c, g, s, q in hits:
            print(f"{c}\t{g}\t{s}\t{q}")
        print(f"\n— {len(hits)}건")
        return

    print(f"총 {len(m.rows)}행 / 고유 (과목,소분류,유형) {len(m.valid3)}개")
    print(f"과목: {m.grades()}")
    cats = {}
    for c, _, _, _ in m.rows:
        cats[c] = cats.get(c, 0) + 1
    print("\n대분류별 유형 수")
    for c, n in sorted(cats.items(), key=lambda x: -x[1]):
        mark = "" if c in CATEGORY_10 else "   ← 확정 10종 목록에 없음"
        print(f"  {c}\t{n}{mark}")
    if m.conflicts:
        print(f"\n(소분류,유형) 하나에 대분류가 둘 이상인 경우 {len(m.conflicts)}건 — 과목까지 포함한 3튜플(cat3)을 쓸 것")
        for key, c1, c2 in m.conflicts[:10]:
            print(f"  {key}: '{c1}' vs '{c2}'")


if __name__ == "__main__":
    main()
